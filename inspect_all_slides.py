from pptx import Presentation
import sys

try:
    sys.stdout.reconfigure(encoding='utf-8')
except AttributeError:
    pass

pptx_path = r"d:\Auto-OnePay\Bai_giang\cowell\Playwright & TypeScript - Session 1 (1).pptx"
prs = Presentation(pptx_path)

print(f"Total Slides: {len(prs.slides)}")
for idx, slide in enumerate(prs.slides):
    print(f"\nSlide {idx+1}: layout='{slide.slide_layout.name}'")
    for j, shape in enumerate(slide.shapes):
        print(f"  Shape {j}: name='{shape.name}', type={shape.shape_type}")
        if shape.has_text_frame and shape.text_frame.text.strip():
            print(f"    Text: '{shape.text_frame.text[:120].strip()}'")
