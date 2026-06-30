import os
import re
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# -----------------------------------------------------------------------------
# CONSTANTS & STYLES (Light Google/Cowell Brand Palette)
# -----------------------------------------------------------------------------
FONT_TITLE = "Open Sans"
FONT_BODY = "Roboto"

TEXT_DARK = RGBColor(32, 33, 36)      # Slate 900 / #202124 (Headers)
TEXT_BODY = RGBColor(60, 64, 67)      # Slate 700 / #3C4043 (Body copy)
TEXT_MUTED = RGBColor(95, 99, 104)    # Slate 500 / #5F6368 (Subtitles/Details)

BLUE_BRAND = RGBColor(26, 115, 232)   # Google Blue / #1A73E8
BLUE_ACCENT = RGBColor(66, 133, 244)  # Google Light Blue Accent / #4285F4
GREEN_BRAND = RGBColor(52, 168, 83)   # Google Green / #34A853
ORANGE_BRAND = RGBColor(249, 171, 0)  # Google Yellow-Orange / #F9AB00
RED_BRAND = RGBColor(234, 67, 53)     # Google Red / #EA4335

CARD_BG = RGBColor(248, 249, 250)     # Off-white Slate / #F8F9FA
CARD_BORDER = RGBColor(218, 220, 224) # Light Border / #DADCE0

# Accent highlights for cards
HIGHLIGHT_BLUE = RGBColor(232, 240, 254)  # Light blue tint / #E8F0FE
HIGHLIGHT_GREEN = RGBColor(230, 244, 234) # Light green tint / #E6F4EA
HIGHLIGHT_ORANGE = RGBColor(254, 247, 224)# Light orange tint / #FEF7E0
HIGHLIGHT_RED = RGBColor(252, 232, 230)   # Light red tint / #FCE8E6

CODE_BG = RGBColor(15, 23, 42)            # Dark slate terminal color / #0F172A
CODE_TEXT = RGBColor(248, 250, 252)       # Light terminal text

# -----------------------------------------------------------------------------
# HELPER FUNCTIONS FOR PRESENTATION STYLING
# -----------------------------------------------------------------------------
def add_slide_header_light(slide, title_text, section_name=None, color=BLUE_ACCENT):
    # Slide main title textbox
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.8))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = FONT_TITLE
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    # Underline horizontal line shape (Google Brand style)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.15), Inches(1.5), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()
    
    # Section Category subtitle (if available)
    if section_name:
        # Strip number like "PHẦN 1: "
        clean_sect = re.sub(r'^PHẦN\s+\d+:\s*', '', section_name).strip()
        txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(1.22), Inches(11.7), Inches(0.4))
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
        p2 = tf2.paragraphs[0]
        p2.text = clean_sect
        p2.font.name = FONT_TITLE
        p2.font.size = Pt(15)
        p2.font.bold = True
        p2.font.color.rgb = BLUE_BRAND

def add_formatted_text_light(paragraph, text, font_size=15, default_color=TEXT_BODY):
    # Strip markdown lists indicators like "- ", "* "
    clean_text = re.sub(r'^[\-\*\+\s]+', '', text).strip()
    
    # Parse bold text **bold**
    parts = re.split(r'(\*\*.*?\*\*)', clean_text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            bold_text = part[2:-2]
            run = paragraph.add_run()
            run.text = bold_text
            run.font.bold = True
            run.font.name = FONT_BODY
            run.font.size = Pt(font_size)
            run.font.color.rgb = BLUE_BRAND
        else:
            if part:
                run = paragraph.add_run()
                run.text = part
                run.font.name = FONT_BODY
                run.font.size = Pt(font_size)
                run.font.color.rgb = default_color

def add_bullet_list_light(slide, items, left, top, width, height, font_size=15, default_color=TEXT_BODY, bullet_color=GREEN_BRAND):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    for i, item in enumerate(items):
        if not item.strip():
            continue
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        p.level = 0
        
        # Add styled checkmark prefix
        run_bullet = p.add_run()
        run_bullet.text = "✓  "
        run_bullet.font.bold = True
        run_bullet.font.name = FONT_BODY
        run_bullet.font.size = Pt(font_size)
        run_bullet.font.color.rgb = bullet_color
        
        add_formatted_text_light(p, item, font_size, default_color)

# -----------------------------------------------------------------------------
# PARSING CODE
# -----------------------------------------------------------------------------
def parse_markdown(filepath):
    with open(filepath, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Split content by sections and slides
    parts = re.split(r'(PHẦN \d+:[^\n]+|SLIDE \d+:[^\n]+)', content)
    
    items = []
    current_section = "TỔNG QUAN KHÓA HỌC"
    
    i = 0
    while i < len(parts):
        part = parts[i].strip()
        if not part:
            i += 1
            continue
            
        if part.startswith("PHẦN"):
            current_section = part
            i += 1
        elif part.startswith("SLIDE"):
            slide_header = part
            slide_body = parts[i+1] if (i+1) < len(parts) else ""
            
            # Extract slide number and title
            slide_match = re.match(r'SLIDE (\d+):\s*(.*)', slide_header)
            slide_num = int(slide_match.group(1)) if slide_match else 0
            slide_title_header = slide_match.group(2).strip() if slide_match else ""
            
            # Parse body
            title, content_text, key_msg, notes = parse_slide_body(slide_body)
            if not title:
                title = slide_title_header
                
            items.append({
                'type': 'slide',
                'section': current_section,
                'num': slide_num,
                'title': title,
                'content': content_text,
                'key_msg': key_msg,
                'notes': notes
            })
            i += 2
        else:
            i += 1
            
    return items

def parse_slide_body(body_text):
    lines = body_text.strip().split('\n')
    current_section = None
    section_data = {
        'title': [],
        'content': [],
        'key_msg': [],
        'notes': []
    }
    
    for line in lines:
        line_strip = line.strip()
        if not line_strip:
            if current_section in ['content', 'notes']:
                section_data[current_section].append("")
            continue
            
        if line_strip.startswith("Tiêu đề:"):
            current_section = "title"
            rest = line_strip[len("Tiêu đề:"):].strip()
            if rest:
                section_data['title'].append(rest)
        elif line_strip.startswith("Nội dung slide:") or line_strip.startswith("Nội dung đề xuất:"):
            current_section = "content"
            rest = line_strip[len("Nội dung slide:"):].strip()
            if rest:
                section_data['content'].append(rest)
        elif line_strip.startswith("Thông điệp chính:") or line_strip.startswith("Thông điệp cuối:"):
            current_section = "key_msg"
            rest = line_strip[len("Thông điệp chính:"):].strip()
            if rest:
                section_data['key_msg'].append(rest)
        elif line_strip.startswith("Speaker note cho trainer:") or line_strip.startswith("Speaker note cho trainer"):
            current_section = "notes"
            rest = line_strip[len("Speaker note cho trainer:"):].strip()
            if rest:
                section_data['notes'].append(rest)
        elif line_strip.startswith("Gợi ý đáp án:") or line_strip.startswith("Lưu ý:") or line_strip.startswith("Kết quả thực tế:") or line_strip.startswith("Kết quả cần đạt:"):
            current_section = "content"
            section_data['content'].append(line_strip)
        elif current_section:
            section_data[current_section].append(line_strip)
        else:
            current_section = "content"
            section_data['content'].append(line_strip)
            
    title = " ".join(section_data['title']).strip()
    content = "\n".join(section_data['content']).strip()
    key_msg = "\n".join(section_data['key_msg']).strip()
    notes = "\n".join(section_data['notes']).strip()
    
    return title, content, key_msg, notes

# -----------------------------------------------------------------------------
# SLIDE GENERATORS FOR LIGHT GOOGLE/COWELL THEME
# -----------------------------------------------------------------------------
def make_title_slide_light(prs, slide_data, title_layout):
    slide = prs.slides.add_slide(title_layout)
    
    # Write to Title placeholder
    title_ph = slide.placeholders[0]
    # Re-structure title with break
    title_text = slide_data['title']
    title_text = title_text.replace(" với ", "\nvới ")
    title_ph.text = title_text
    
    p_title = title_ph.text_frame.paragraphs[0]
    p_title.font.name = FONT_TITLE
    p_title.font.size = Pt(40)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_DARK
    
    # Write to Subtitle placeholder
    subtitle_ph = slide.placeholders[1]
    tf = subtitle_ph.text_frame
    tf.word_wrap = True
    tf.text = "Buổi 1: Tổng quan & Khởi động"
    
    # Subtitle formatting
    p_sub = tf.paragraphs[0]
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = TEXT_MUTED
    
    p_course = tf.add_paragraph()
    p_course.text = "Khóa học: Automation Testing với Playwright & TypeScript"
    p_course.font.name = FONT_BODY
    p_course.font.size = Pt(15)
    p_course.font.bold = True
    p_course.font.color.rgb = BLUE_BRAND
    p_course.space_before = Pt(12)
    
    # Add minor details at the bottom of slide if any content exists
    # Find list items from slide content
    items = [line for line in slide_data['content'].split('\n') if line.strip()]
    if items:
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(4.8), Inches(11.7), Inches(1.8))
        tf_d = txBox.text_frame
        tf_d.word_wrap = True
        for i, item in enumerate(items):
            p = tf_d.paragraphs[0] if i == 0 else tf_d.add_paragraph()
            p.space_after = Pt(4)
            add_formatted_text_light(p, item, font_size=13, default_color=TEXT_MUTED)

    if slide_data['notes']:
        slide.notes_slide.notes_text_frame.text = slide_data['notes']

def make_section_divider_light(prs, section_title, section_layout):
    slide = prs.slides.add_slide(section_layout)
    
    # E.g. "PHẦN 1: MỞ ĐẦU KHÓA HỌC"
    match = re.match(r'(PHẦN \d+):\s*(.*)', section_title)
    if match:
        part_no = match.group(1)
        part_name = match.group(2)
        
        title_ph = slide.placeholders[0]
        title_ph.text = part_no.upper()
        p1 = title_ph.text_frame.paragraphs[0]
        p1.font.name = FONT_TITLE
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = BLUE_BRAND
        
        body_ph = slide.placeholders[1]
        body_ph.text = part_name.upper()
        p2 = body_ph.text_frame.paragraphs[0]
        p2.font.name = FONT_TITLE
        p2.font.size = Pt(36)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_DARK
    else:
        title_ph = slide.placeholders[0]
        title_ph.text = "CHƯƠNG"
        body_ph = slide.placeholders[1]
        body_ph.text = section_title.upper()

def make_card_layout_light(prs, slide_data, num_cards, blank_layout):
    slide = prs.slides.add_slide(blank_layout)
    add_slide_header_light(slide, slide_data['title'], slide_data['section'])
    
    lines = [line.strip() for line in slide_data['content'].split('\n') if line.strip()]
    
    cards = []
    current_card = None
    for line in lines:
        if re.match(r'^\d+\.', line) or line.startswith('- **') or line.startswith('**'):
            if current_card:
                cards.append(current_card)
            current_card = {'title': line, 'bullets': []}
        else:
            if current_card:
                current_card['bullets'].append(line)
            else:
                current_card = {'title': "Details", 'bullets': [line]}
    if current_card:
        cards.append(current_card)
        
    if len(cards) < num_cards and len(lines) >= num_cards:
        cards = []
        for i in range(num_cards):
            cards.append({'title': lines[i], 'bullets': []})
            
    cards = cards[:num_cards]
    
    total_width = Inches(11.7)
    gap = Inches(0.3)
    card_width = (total_width - (gap * (num_cards - 1))) / num_cards
    left_start = Inches(0.8)
    top_pos = Inches(1.9)
    card_height = Inches(4.5)
    
    for idx, card in enumerate(cards):
        left_pos = left_start + idx * (card_width + gap)
        
        # Rounded rectangle shape
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, card_width, card_height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = CARD_BORDER
        shape.line.width = Pt(1.5)
        
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.18)
        tf.margin_right = Inches(0.18)
        tf.margin_top = Inches(0.18)
        tf.margin_bottom = Inches(0.18)
        
        # Card Header
        p = tf.paragraphs[0]
        title_text = re.sub(r'^\d+\.\s*', '', card['title'])
        title_text = re.sub(r'^\*\*|\*\*$', '', title_text)
        p.text = title_text.upper()
        p.font.name = FONT_TITLE
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = BLUE_BRAND
        p.space_after = Pt(12)
        p.alignment = PP_ALIGN.CENTER
        
        # Bullets
        for bullet in card['bullets']:
            p_bullet = tf.add_paragraph()
            p_bullet.space_after = Pt(6)
            
            run_bullet = p_bullet.add_run()
            run_bullet.text = "✓  "
            run_bullet.font.bold = True
            run_bullet.font.color.rgb = GREEN_BRAND
            
            add_formatted_text_light(p_bullet, bullet, font_size=13)
            
    # Add key message if exists
    if slide_data['key_msg']:
        tx = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.5))
        tf = tx.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_formatted_text_light(p, "**Thông điệp:** " + slide_data['key_msg'], font_size=14, default_color=TEXT_MUTED)
        
    if slide_data['notes']:
        slide.notes_slide.notes_text_frame.text = slide_data['notes']

def make_two_column_layout_light(prs, slide_data, column_titles=None, left_accent=BLUE_BRAND, right_accent=BLUE_BRAND, blank_layout=None):
    slide = prs.slides.add_slide(blank_layout)
    add_slide_header_light(slide, slide_data['title'], slide_data['section'])
    
    lines = [line.strip() for line in slide_data['content'].split('\n') if line.strip()]
    
    left_items = []
    right_items = []
    
    split_words = ["Nhưng Automation không thay thế được:", "Nhưng không thay thế:", "Không thay thế:", "Automation KHÔNG THỂ thay thế:", "Automation test:", "Expected result", "Kết quả thực tế:", "Bài tập 3:"]
    split_idx = len(lines) // 2
    
    for i, line in enumerate(lines):
        if any(word in line for word in split_words):
            split_idx = i
            break
            
    left_items = lines[:split_idx]
    right_items = lines[split_idx:]
    
    col_width = Inches(5.6)
    left_col_x = Inches(0.8)
    right_col_x = Inches(6.9)
    top_pos = Inches(1.9)
    col_height = Inches(4.5)
    
    # Left Box
    shape_l = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_col_x, top_pos, col_width, col_height)
    shape_l.fill.solid()
    # Apply soft color highlights if specific colors used
    if left_accent == RED_BRAND:
        shape_l.fill.fore_color.rgb = HIGHLIGHT_RED
    elif left_accent == GREEN_BRAND:
        shape_l.fill.fore_color.rgb = HIGHLIGHT_GREEN
    else:
        shape_l.fill.fore_color.rgb = CARD_BG
        
    shape_l.line.color.rgb = left_accent
    shape_l.line.width = Pt(1.5)
    tf_l = shape_l.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = tf_l.margin_right = tf_l.margin_top = tf_l.margin_bottom = Inches(0.2)
    
    if column_titles and len(column_titles) >= 1:
        p = tf_l.paragraphs[0]
        p.text = column_titles[0].upper()
        p.font.name = FONT_TITLE
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = left_accent
        p.space_after = Pt(10)
        start_p_idx = 1
    else:
        start_p_idx = 0
        
    for i, item in enumerate(left_items):
        p_item = tf_l.add_paragraph() if (i > 0 or start_p_idx > 0) else tf_l.paragraphs[0]
        p_item.space_after = Pt(6)
        
        run_bullet = p_item.add_run()
        run_bullet.text = "✓  "
        run_bullet.font.bold = True
        run_bullet.font.color.rgb = left_accent
        
        add_formatted_text_light(p_item, item, font_size=13)
        
    # Right Box
    shape_r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_col_x, top_pos, col_width, col_height)
    shape_r.fill.solid()
    if right_accent == GREEN_BRAND:
        shape_r.fill.fore_color.rgb = HIGHLIGHT_GREEN
    elif right_accent == ORANGE_BRAND:
        shape_r.fill.fore_color.rgb = HIGHLIGHT_ORANGE
    else:
        shape_r.fill.fore_color.rgb = CARD_BG
        
    shape_r.line.color.rgb = right_accent
    shape_r.line.width = Pt(1.5)
    tf_r = shape_r.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_right = tf_r.margin_top = tf_r.margin_bottom = Inches(0.2)
    
    if column_titles and len(column_titles) >= 2:
        p = tf_r.paragraphs[0]
        p.text = column_titles[1].upper()
        p.font.name = FONT_TITLE
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = right_accent
        p.space_after = Pt(10)
        start_p_idx = 1
    else:
        start_p_idx = 0
        
    for i, item in enumerate(right_items):
        p_item = tf_r.add_paragraph() if (i > 0 or start_p_idx > 0) else tf_r.paragraphs[0]
        p_item.space_after = Pt(6)
        
        run_bullet = p_item.add_run()
        run_bullet.text = "✓  "
        run_bullet.font.bold = True
        run_bullet.font.color.rgb = right_accent
        
        add_formatted_text_light(p_item, item, font_size=13)
        
    # Key message at the bottom
    if slide_data['key_msg']:
        tx = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.5))
        tf = tx.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_formatted_text_light(p, "**Thông điệp:** " + slide_data['key_msg'], font_size=14, default_color=TEXT_MUTED)
        
    if slide_data['notes']:
        slide.notes_slide.notes_text_frame.text = slide_data['notes']

def make_pyramid_slide_light(prs, slide_data, blank_layout):
    slide = prs.slides.add_slide(blank_layout)
    add_slide_header_light(slide, slide_data['title'], slide_data['section'])
    
    pyramid_center_x = Inches(3.2)
    
    # 3. E2E Test (Top Tier) - Google Blue background
    shape1 = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, pyramid_center_x - Inches(1.5), Inches(2.2), Inches(3.0), Inches(1.2))
    shape1.fill.solid()
    shape1.fill.fore_color.rgb = BLUE_BRAND
    shape1.line.color.rgb = CARD_BORDER
    shape1.line.width = Pt(1.5)
    tf1 = shape1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    p1.text = "E2E Test\n(Playwright)"
    p1.font.name = FONT_TITLE
    p1.font.bold = True
    p1.font.size = Pt(14)
    p1.font.color.rgb = RGBColor(255, 255, 255)
    
    # 2. Integration Test (Middle Tier) - Light Blue background
    shape2 = slide.shapes.add_shape(MSO_SHAPE.TRAPEZOID, pyramid_center_x - Inches(2.5), Inches(3.5), Inches(5.0), Inches(1.2))
    shape2.fill.solid()
    shape2.fill.fore_color.rgb = HIGHLIGHT_BLUE
    shape2.line.color.rgb = CARD_BORDER
    shape2.line.width = Pt(1.5)
    tf2 = shape2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "Integration / API Test"
    p2.font.name = FONT_TITLE
    p2.font.bold = True
    p2.font.size = Pt(14)
    p2.font.color.rgb = BLUE_BRAND
    
    # 1. Unit Test (Bottom Tier) - Off-white card background
    shape3 = slide.shapes.add_shape(MSO_SHAPE.TRAPEZOID, pyramid_center_x - Inches(3.5), Inches(4.8), Inches(7.0), Inches(1.2))
    shape3.fill.solid()
    shape3.fill.fore_color.rgb = CARD_BG
    shape3.line.color.rgb = CARD_BORDER
    shape3.line.width = Pt(1.5)
    tf3 = shape3.text_frame
    tf3.word_wrap = True
    p3 = tf3.paragraphs[0]
    p3.alignment = PP_ALIGN.CENTER
    p3.text = "Unit Test\n(Developer-written)"
    p3.font.name = FONT_TITLE
    p3.font.bold = True
    p3.font.size = Pt(14)
    p3.font.color.rgb = TEXT_BODY
    
    # Right column text details
    right_x = Inches(7.0)
    width = Inches(5.5)
    items = [line for line in slide_data['content'].split('\n') if line.strip()]
    add_bullet_list_light(slide, items, right_x, Inches(2.0), width, Inches(4.5), font_size=14)
    
    if slide_data['notes']:
        slide.notes_slide.notes_text_frame.text = slide_data['notes']

def make_roadmap_slide_light(prs, slide_data, num_steps, blank_layout):
    slide = prs.slides.add_slide(blank_layout)
    add_slide_header_light(slide, slide_data['title'], slide_data['section'])
    
    items = [line.strip() for line in slide_data['content'].split('\n') if line.strip()]
    
    steps = []
    current_step = None
    for item in items:
        if re.match(r'^\d+\.', item) or item.startswith('- '):
            if current_step:
                steps.append(current_step)
            current_step = {'title': item, 'desc': []}
        else:
            if current_step:
                current_step['desc'].append(item)
            else:
                current_step = {'title': "Step", 'desc': [item]}
    if current_step:
        steps.append(current_step)
        
    steps = steps[:num_steps]
    if len(steps) < num_steps:
        steps = [{'title': line, 'desc': []} for line in items[:num_steps]]
        
    top_start = Inches(1.9)
    left_pos = Inches(0.8)
    width = Inches(11.7)
    height = Inches(0.85)
    gap = Inches(0.12)
    
    for idx, step in enumerate(steps):
        top_pos = top_start + idx * (height + gap)
        
        # Step container card
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = CARD_BORDER
        shape.line.width = Pt(1)
        
        # Number badge shape
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, left_pos + Inches(0.15), top_pos + Inches(0.15), Inches(0.55), Inches(0.55))
        badge.fill.solid()
        badge.fill.fore_color.rgb = BLUE_BRAND
        badge.line.fill.background()
        tf_b = badge.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
        p_b = tf_b.paragraphs[0]
        p_b.alignment = PP_ALIGN.CENTER
        p_b.text = f"{idx+1:02d}"
        p_b.font.name = FONT_TITLE
        p_b.font.size = Pt(16)
        p_b.font.bold = True
        p_b.font.color.rgb = RGBColor(255, 255, 255)
        
        # Text details
        tx = slide.shapes.add_textbox(left_pos + Inches(0.9), top_pos + Inches(0.1), width - Inches(1.1), height - Inches(0.2))
        tf = tx.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        title_text = re.sub(r'^\d+\.\s*', '', step['title'])
        title_text = re.sub(r'^[\-\*\+\s]+', '', title_text)
        title_text = re.sub(r'^\*\*|\*\*$', '', title_text)
        p.text = title_text
        p.font.name = FONT_TITLE
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        
        if step['desc']:
            p2 = tf.add_paragraph()
            p2.text = " — ".join(step['desc'])
            p2.font.name = FONT_BODY
            p2.font.size = Pt(12)
            p2.font.color.rgb = TEXT_MUTED
            p2.space_before = Pt(2)
            
    if slide_data['notes']:
        slide.notes_slide.notes_text_frame.text = slide_data['notes']

def make_modules_grid_slide_light(prs, slide_data, blank_layout):
    slide = prs.slides.add_slide(blank_layout)
    add_slide_header_light(slide, slide_data['title'], slide_data['section'])
    
    modules = [
        {"num": "M1", "title": "Course Intro", "desc": "Tổng quan Automation Testing, Playwright & TypeScript."},
        {"num": "M2", "title": "TypeScript Basic", "desc": "Học TypeScript ở mức đủ dùng cho QA viết automation."},
        {"num": "M3", "title": "Playwright Basic", "desc": "Viết test UI cơ bản, locator, action & assertion."},
        {"num": "M4", "title": "Playwright Advanced", "desc": "Page Object Model, Fixture, Data-driven & AI bonus."},
        {"num": "M5", "title": "Analysis & Report", "desc": "Debug, HTML Report, Trace Viewer & flaky test."},
        {"num": "M6", "title": "Capstone & Interview", "desc": "Hoàn thiện project cuối khóa, tối ưu CV, phỏng vấn."}
    ]
    
    card_width = Inches(3.7)
    card_height = Inches(2.2)
    gap_x = Inches(0.3)
    gap_y = Inches(0.3)
    left_start = Inches(0.8)
    top_start = Inches(1.9)
    
    for idx, mod in enumerate(modules):
        row = idx // 3
        col = idx % 3
        
        left_pos = left_start + col * (card_width + gap_x)
        top_pos = top_start + row * (card_height + gap_y)
        
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, card_width, card_height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = CARD_BORDER
        shape.line.width = Pt(1.5)
        
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(0.15)
        
        p_badge = tf.paragraphs[0]
        p_badge.text = mod['num']
        p_badge.font.name = FONT_TITLE
        p_badge.font.size = Pt(12)
        p_badge.font.bold = True
        p_badge.font.color.rgb = BLUE_BRAND
        p_badge.space_after = Pt(4)
        
        p_title = tf.add_paragraph()
        p_title.text = mod['title'].upper()
        p_title.font.name = FONT_TITLE
        p_title.font.size = Pt(15)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_DARK
        p_title.space_after = Pt(8)
        
        p_desc = tf.add_paragraph()
        p_desc.text = mod['desc']
        p_desc.font.name = FONT_BODY
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = TEXT_BODY
        
    if slide_data['notes']:
        slide.notes_slide.notes_text_frame.text = slide_data['notes']

def make_code_demo_slide_light(prs, slide_data, blank_layout):
    slide = prs.slides.add_slide(blank_layout)
    add_slide_header_light(slide, slide_data['title'], slide_data['section'])
    
    # Left Column
    left_x = Inches(0.8)
    col_width = Inches(5.0)
    items = [
        "Mở trang Web (https://demo.playwright.dev/todomvc)",
        "Nhập một Todo Item vào ô textbox",
        "Nhấn phím Enter để thêm mới",
        "Kiểm tra item vừa thêm có hiển thị trên danh sách",
        "Playwright tự động ghi nhận pass/fail, xuất report"
    ]
    add_bullet_list_light(slide, items, left_x, Inches(2.0), col_width, Inches(4.5), font_size=15)
    
    # Right Column: Code block (still dark, representing the editor)
    right_x = Inches(6.2)
    code_width = Inches(6.3)
    code_height = Inches(4.5)
    
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x, Inches(2.0), code_width, code_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = CODE_BG
    shape.line.color.rgb = BLUE_BRAND
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.25)
    
    p_hdr = tf.paragraphs[0]
    p_hdr.text = "/// demo.spec.ts (TypeScript)"
    p_hdr.font.name = "Consolas"
    p_hdr.font.size = Pt(12)
    p_hdr.font.bold = True
    p_hdr.font.color.rgb = TEXT_MUTED
    p_hdr.space_after = Pt(15)
    
    code_lines = [
        "import { test, expect } from '@playwright/test';",
        "",
        "test('user can add a todo item', async ({ page }) => {",
        "  await page.goto('https://demo.playwright.dev/todomvc');",
        "",
        "  await page.getByPlaceholder('What needs to be done?')",
        "            .fill('Learn Playwright');",
        "  await page.keyboard.press('Enter');",
        "",
        "  await expect(page.getByText('Learn Playwright'))",
        "        .toBeVisible();",
        "});"
    ]
    
    for cl in code_lines:
        p = tf.add_paragraph()
        p.text = "" # clear text to insert formatted runs
        words = re.split(r'(\bimport\b|\bfrom\b|\btest\b|\bexpect\b|\basync\b|\bawait\b|\'Learn Playwright\'|\'https://demo.playwright.dev/todomvc\')', cl)
        for w in words:
            r = p.add_run()
            r.text = w
            r.font.name = "Consolas"
            r.font.size = Pt(11.5)
            if w in ['import', 'from', 'test', 'expect', 'async', 'await']:
                r.font.color.rgb = RGBColor(56, 189, 248) # Cyan in terminal
                r.font.bold = True
            elif w.startswith("'") or w.startswith('"'):
                r.font.color.rgb = RGBColor(245, 158, 11) # Orange in terminal
            else:
                r.font.color.rgb = CODE_TEXT
                
    if slide_data['notes']:
        slide.notes_slide.notes_text_frame.text = slide_data['notes']

def make_ai_support_slide_light(prs, slide_data, blank_layout):
    slide = prs.slides.add_slide(blank_layout)
    add_slide_header_light(slide, slide_data['title'], slide_data['section'])
    
    # Left
    left_x = Inches(0.8)
    col_width = Inches(5.2)
    items = [
        "**AI là trợ lý (Assistant)**: Hỗ trợ học nhanh hơn khi gặp khái niệm khó.",
        "**Giải thích code**: Giúp QA no-tech đọc hiểu code Playwright và TypeScript.",
        "**Tối ưu hóa**: Gợi ý cách viết script gọn và dễ bảo trì.",
        "**Hỗ trợ Debug**: Giải thích log lỗi chi tiết khi test fail.",
        "**Sinh code mẫu**: Tạo script khung từ requirement hoặc checklist manual.",
        "**LƯU Ý**: Không copy máy móc! QA cần tự review code và expected result."
    ]
    add_bullet_list_light(slide, items, left_x, Inches(2.0), col_width, Inches(4.5), font_size=14)
    
    # Right: AI Prompt template block
    right_x = Inches(6.4)
    code_width = Inches(6.1)
    code_height = Inches(4.5)
    
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x, Inches(2.0), code_width, code_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = CODE_BG
    shape.line.color.rgb = ORANGE_BRAND
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.25)
    
    p_hdr = tf.paragraphs[0]
    p_hdr.text = "/// PROMPT AI THAM KHẢO"
    p_hdr.font.name = FONT_TITLE
    p_hdr.font.size = Pt(13)
    p_hdr.font.bold = True
    p_hdr.font.color.rgb = ORANGE_BRAND
    p_hdr.space_after = Pt(15)
    
    prompt_text = (
        "\"Đóng vai trò là Automation QA Trainer.\n"
        "Hãy giải thích đoạn code Playwright sau cho người mới bắt đầu, chưa biết nhiều về code.\n"
        "Giải thích từng dòng bằng ngôn ngữ đơn giản và cho biết dòng nào là action, dòng nào là assertion.\n\n"
        "[Paste code Playwright của bạn vào đây]\""
    )
    
    p_body = tf.add_paragraph()
    p_body.text = prompt_text
    p_body.font.name = "Consolas"
    p_body.font.size = Pt(12)
    p_body.font.color.rgb = CODE_TEXT
    p_body.font.italic = True
    
    if slide_data['notes']:
        slide.notes_slide.notes_text_frame.text = slide_data['notes']

def make_checklist_slide_light(prs, slide_data, blank_layout):
    slide = prs.slides.add_slide(blank_layout)
    add_slide_header_light(slide, slide_data['title'], slide_data['section'])
    
    lines = [line.strip() for line in slide_data['content'].split('\n') if line.strip()]
    
    split_idx = (len(lines) + 1) // 2
    left_lines = lines[:split_idx]
    right_lines = lines[split_idx:]
    
    left_x = Inches(0.8)
    right_x = Inches(6.8)
    width = Inches(5.7)
    top_pos = Inches(1.9)
    
    def render_column(slide, items, col_left):
        txBox = slide.shapes.add_textbox(col_left, top_pos, width, Inches(5.0))
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(8)
            p.level = 0
            
            # Checkbox symbol ☑ in green/blue
            run_check = p.add_run()
            run_check.text = "☑  "
            run_check.font.name = FONT_BODY
            run_check.font.size = Pt(14)
            run_check.font.bold = True
            run_check.font.color.rgb = BLUE_BRAND
            
            clean_item = re.sub(r'^\[[\sx]\]\s*', '', item)
            add_formatted_text_light(p, clean_item, font_size=13)
            
    render_column(slide, left_lines, left_x)
    render_column(slide, right_lines, right_x)
    
    if slide_data['notes']:
        slide.notes_slide.notes_text_frame.text = slide_data['notes']

def make_activity_slide_light(prs, slide_data, blank_layout):
    slide = prs.slides.add_slide(blank_layout)
    add_slide_header_light(slide, slide_data['title'], slide_data['section'], color=ORANGE_BRAND)
    
    lines = [line.strip() for line in slide_data['content'].split('\n') if line.strip()]
    
    left_x = Inches(0.8)
    col_width = Inches(5.4)
    
    right_x = Inches(6.6)
    card_width = Inches(5.9)
    card_height = Inches(4.5)
    
    q_lines = []
    d_lines = []
    current_list = q_lines
    for line in lines:
        if "Câu hỏi" in line or "thảo luận" in line:
            current_list = d_lines
        current_list.append(line)
        
    add_bullet_list_light(slide, q_lines, left_x, Inches(2.0), col_width, Inches(4.5), font_size=15, bullet_color=ORANGE_BRAND)
    
    # Discussion box in highlights orange
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x, Inches(2.0), card_width, card_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = HIGHLIGHT_ORANGE
    shape.line.color.rgb = ORANGE_BRAND
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.2)
    
    p = tf.paragraphs[0]
    p.text = "CÂU HỎI THẢO LUẬN NHÓM"
    p.font.name = FONT_TITLE
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ORANGE_BRAND
    p.space_after = Pt(12)
    p.alignment = PP_ALIGN.CENTER
    
    if len(d_lines) <= 1:
        d_lines = [
            "Công việc nào làm bạn mất nhiều thời gian nhất?",
            "Công việc nào dễ bị sai sót khi làm thủ công?",
            "Công việc nào nếu tự động hóa được sẽ giúp tiết kiệm nhiều thời gian?"
        ]
        
    for item in d_lines:
        if "Câu hỏi thảo luận" in item:
            continue
        p_item = tf.add_paragraph()
        p_item.space_after = Pt(10)
        
        run_q = p_item.add_run()
        run_q.text = "❓  "
        
        add_formatted_text_light(p_item, item, font_size=14)
        
    if slide_data['notes']:
        slide.notes_slide.notes_text_frame.text = slide_data['notes']

def make_quiz_slide_light(prs, slide_data, blank_layout):
    slide = prs.slides.add_slide(blank_layout)
    add_slide_header_light(slide, slide_data['title'], slide_data['section'], color=ORANGE_BRAND)
    
    content = slide_data['content']
    questions = []
    
    lines = content.split('\n')
    for line in lines:
        l = line.strip()
        if re.match(r'^\d+\.', l):
            questions.append(l)
            
    if len(questions) == 0:
        questions = [l for l in lines if l.strip()][:8]
        
    split_idx = (len(questions) + 1) // 2
    left_qs = questions[:split_idx]
    right_qs = questions[split_idx:]
    
    left_x = Inches(0.8)
    right_x = Inches(6.8)
    width = Inches(5.7)
    top_pos = Inches(1.9)
    
    def render_quiz_column(slide, items, col_left):
        txBox = slide.shapes.add_textbox(col_left, top_pos, width, Inches(5.0))
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(12)
            
            num_match = re.match(r'^(\d+)\.\s*(.*)', item)
            if num_match:
                num = num_match.group(1)
                text = num_match.group(2)
            else:
                num = str(i+1)
                text = item
                
            run_num = p.add_run()
            run_num.text = f"{num}.  "
            run_num.font.bold = True
            run_num.font.name = FONT_TITLE
            run_num.font.size = Pt(14)
            run_num.font.color.rgb = ORANGE_BRAND
            
            add_formatted_text_light(p, text, font_size=13)
            
    render_quiz_column(slide, left_qs, left_x)
    render_quiz_column(slide, right_qs, right_x)
    
    if slide_data['notes']:
        slide.notes_slide.notes_text_frame.text = slide_data['notes']

def make_standard_slide_light(prs, slide_data, blank_layout):
    slide = prs.slides.add_slide(blank_layout)
    add_slide_header_light(slide, slide_data['title'], slide_data['section'])
    
    lines = [line.strip() for line in slide_data['content'].split('\n') if line.strip()]
    font_size = 16 if len(lines) < 6 else 14
    
    content_height = Inches(4.2)
    if slide_data['key_msg']:
        content_height = Inches(3.6)
        
    add_bullet_list_light(slide, lines, Inches(0.8), Inches(1.9), Inches(11.7), content_height, font_size=font_size)
    
    if slide_data['key_msg']:
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.5), Inches(11.7), Inches(1.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = HIGHLIGHT_BLUE
        shape.line.color.rgb = BLUE_BRAND
        shape.line.width = Pt(1.5)
        
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.15)
        p = tf.paragraphs[0]
        p.text = "KEY TAKEAWAY"
        p.font.name = FONT_TITLE
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = BLUE_BRAND
        p.space_after = Pt(4)
        
        p2 = tf.add_paragraph()
        add_formatted_text_light(p2, slide_data['key_msg'], font_size=13)
        
    if slide_data['notes']:
        slide.notes_slide.notes_text_frame.text = slide_data['notes']

# -----------------------------------------------------------------------------
# MAIN PROCESS
# -----------------------------------------------------------------------------
def main():
    import sys
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except AttributeError:
        pass
        
    md_path = r"d:\Auto-OnePay\Bai_giang\cowell\Noi_dung_Buoi_1_Playwright_Bo_sung_Module_AI.md"
    template_path = r"d:\Auto-OnePay\Bai_giang\cowell\Playwright & TypeScript - Session 1 (1).pptx"
    pptx_path = r"d:\Auto-OnePay\Bai_giang\cowell\Noi_dung_Buoi_1_Playwright_Bo_sung_Module_AI.pptx"
    
    print(f"Reading and parsing markdown: {md_path}")
    if not os.path.exists(md_path):
        print("Error: Markdown file not found!")
        return
        
    slides_data = parse_markdown(md_path)
    print(f"Parsed {len(slides_data)} slides successfully.")
    
    print(f"Loading reference template presentation: {template_path}")
    if not os.path.exists(template_path):
        print("Error: Reference template presentation file not found!")
        return
        
    prs = Presentation(template_path)
    
    # Map slide layouts
    layout_map = {l.name: l for l in prs.slide_layouts}
    title_layout = layout_map.get('TITLE', prs.slide_layouts[1])
    section_layout = layout_map.get('SECTION_HEADER', prs.slide_layouts[3])
    blank_layout = layout_map.get('BLANK', prs.slide_layouts[0])
    
    # Delete existing slides from template
    for i in range(len(prs.slides)-1, -1, -1):
        rId = prs.slides._sldIdLst[i].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[i]
        
    last_section = None
    
    for idx, slide_data in enumerate(slides_data):
        current_section = slide_data['section']
        
        # Section divider slide
        if current_section != last_section:
            print(f"Generating Section Divider: {current_section.encode('ascii', errors='ignore').decode()}")
            make_section_divider_light(prs, current_section, section_layout)
            last_section = current_section
            
        print(f"Generating Slide {slide_data['num']}: {slide_data['title'].encode('ascii', errors='ignore').decode()}")
        
        num = slide_data['num']
        
        if num == 1:
            make_title_slide_light(prs, slide_data, title_layout)
        elif num == 2:
            make_two_column_layout_light(prs, slide_data, column_titles=["Các công việc lặp lại tẻ nhạt...", "NỖI ĐAU THỰC TẾ"], left_accent=RED_BRAND, right_accent=ORANGE_BRAND, blank_layout=blank_layout)
        elif num == 3:
            make_card_layout_light(prs, slide_data, num_cards=3, blank_layout=blank_layout)
        elif num == 4:
            make_two_column_layout_light(prs, slide_data, column_titles=["Những gì Automation làm cực tốt", "Những gì Manual QA làm tốt hơn"], left_accent=GREEN_BRAND, right_accent=BLUE_BRAND, blank_layout=blank_layout)
        elif num == 5:
            make_roadmap_slide_light(prs, slide_data, num_steps=5, blank_layout=blank_layout)
        elif num == 7:
            make_pyramid_slide_light(prs, slide_data, blank_layout=blank_layout)
        elif num == 8:
            make_two_column_layout_light(prs, slide_data, column_titles=["Playwright có thể làm gì?", "Hỗ trợ Trình Duyệt chính"], left_accent=BLUE_BRAND, right_accent=BLUE_BRAND, blank_layout=blank_layout)
        elif num == 10:
            make_two_column_layout_light(prs, slide_data, column_titles=["Vì sao dùng TypeScript?", "Cảnh báo & Gợi ý từ IDE"], left_accent=BLUE_BRAND, right_accent=ORANGE_BRAND, blank_layout=blank_layout)
        elif num == 11:
            make_card_layout_light(prs, slide_data, num_cards=4, blank_layout=blank_layout)
        elif num == 12:
            make_code_demo_slide_light(prs, slide_data, blank_layout=blank_layout)
        elif num == 13:
            make_two_column_layout_light(prs, slide_data, column_titles=["Các thao tác Manual QA", "Lệnh Playwright tương ứng"], left_accent=BLUE_BRAND, right_accent=GREEN_BRAND, blank_layout=blank_layout)
        elif num == 14:
            make_modules_grid_slide_light(prs, slide_data, blank_layout=blank_layout)
        elif num == 21:
            make_two_column_layout_light(prs, slide_data, column_titles=["Giai đoạn Khởi động & Nền tảng (1-5)", "Giai đoạn Thực hành & Project (6-10)"], left_accent=BLUE_BRAND, right_accent=BLUE_BRAND, blank_layout=blank_layout)
        elif num == 22:
            make_two_column_layout_light(prs, slide_data, column_titles=["Kỹ năng thực hành đạt được", "Kết quả thực tế trong CV"], left_accent=GREEN_BRAND, right_accent=BLUE_BRAND, blank_layout=blank_layout)
        elif num == 23:
            make_ai_support_slide_light(prs, slide_data, blank_layout=blank_layout)
        elif num == 24:
            make_two_column_layout_light(prs, slide_data, column_titles=["Học viên chuẩn bị trước Buổi 2", "Ghi chú cài đặt phần mềm"], left_accent=BLUE_BRAND, right_accent=ORANGE_BRAND, blank_layout=blank_layout)
        elif num == 25:
            make_activity_slide_light(prs, slide_data, blank_layout=blank_layout)
        elif num == 26:
            make_quiz_slide_light(prs, slide_data, blank_layout=blank_layout)
        elif num == 27:
            make_two_column_layout_light(prs, slide_data, column_titles=["Bài tập về nhà 1 & 2", "Bài tập 3: Chuẩn bị máy tính"], left_accent=BLUE_BRAND, right_accent=ORANGE_BRAND, blank_layout=blank_layout)
        elif num == 28:
            make_checklist_slide_light(prs, slide_data, blank_layout=blank_layout)
        else:
            make_standard_slide_light(prs, slide_data, blank_layout=blank_layout)
            
    print(f"Saving new presentation to: {pptx_path}")
    prs.save(pptx_path)
    print("Done!")

if __name__ == '__main__':
    main()
