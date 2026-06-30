from pptx import Presentation
from pptx.dml.color import RGBColor
import sys

try:
    sys.stdout.reconfigure(encoding='utf-8')
except AttributeError:
    pass

pptx_path = r"d:\Auto-OnePay\Bai_giang\cowell\Playwright & TypeScript - Session 1 (1).pptx"
prs = Presentation(pptx_path)

print(f"Number of slides in template: {len(prs.slides)}")

def print_font_details(run):
    font = run.font
    color_type = "None"
    color_val = "None"
    if font.color and font.color.type is not None:
        color_type = str(font.color.type)
        try:
            color_val = str(font.color.rgb)
        except Exception:
            color_val = "Theme color or other"
            
    print(f"    Text: '{run.text}'")
    print(f"    Font: Name={font.name}, Size={font.size}, Bold={font.bold}, ColorType={color_type}, Color={color_val}")

for i, slide in enumerate(prs.slides):
    if i >= 5:
        break
    print(f"\n--- Slide {i+1} ---")
    
    # Background
    bg = slide.background
    if bg and bg.fill:
        print(f"  Background type: {bg.fill.type}")
        if bg.fill.type == 1: # Solid
            print(f"  Background Color: {bg.fill.fore_color.rgb if bg.fill.fore_color else 'None'}")
            
    # Shapes
    print(f"  Number of shapes: {len(slide.shapes)}")
    for j, shape in enumerate(slide.shapes):
        print(f"  Shape {j}: '{shape.name}', Type: {shape.shape_type}")
        if shape.has_text_frame:
            tf = shape.text_frame
            print(f"    Text frame text: '{tf.text[:100]}...'")
            for paragraph in tf.paragraphs[:2]:
                for run in paragraph.runs[:2]:
                    print_font_details(run)
                    
        # Check solid fill of shape
        if hasattr(shape, 'fill') and shape.fill and shape.fill.type == 1: # Solid
            print(f"    Fill color: {shape.fill.fore_color.rgb if shape.fill.fore_color else 'None'}")
        if hasattr(shape, 'line') and shape.line and shape.line.color:
            try:
                print(f"    Line color: {shape.line.color.rgb}")
            except Exception:
                pass
