from pptx import Presentation
import sys

try:
    sys.stdout.reconfigure(encoding='utf-8')
except AttributeError:
    pass

pptx_path = r"d:\Auto-OnePay\Bai_giang\cowell\Playwright & TypeScript - Session 1 (1).pptx"
prs = Presentation(pptx_path)

for idx, slide in enumerate(prs.slides):
    print(f"\nSlide {idx+1}:")
    for j, shape in enumerate(slide.shapes):
        if shape.shape_type == 13: # PICTURE
            left_in = shape.left / 914400
            top_in = shape.top / 914400
            width_in = shape.width / 914400
            height_in = shape.height / 914400
            print(f"  Picture shape {j}: left={left_in:.2f} in, top={top_in:.2f} in, width={width_in:.2f} in, height={height_in:.2f} in")
