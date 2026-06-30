from pptx import Presentation

pptx_path = r"d:\Auto-OnePay\Bai_giang\cowell\Playwright & TypeScript - Session 1 (1).pptx"
prs = Presentation(pptx_path)

print(f"Slide width: {prs.slide_width / 914400} inches")
print(f"Slide height: {prs.slide_height / 914400} inches")
print(f"Number of slide layouts: {len(prs.slide_layouts)}")
for i, l in enumerate(prs.slide_layouts):
    print(f"Layout {i}: Name='{l.name}'")
    # Let's inspect placeholders in each layout
    print("  Placeholders:")
    for ph in l.placeholders:
        print(f"    idx={ph.placeholder_format.idx}, name='{ph.name}', type={ph.placeholder_format.type}")
